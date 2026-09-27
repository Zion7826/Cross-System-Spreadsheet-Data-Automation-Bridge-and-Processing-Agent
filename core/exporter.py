# -*- coding: utf-8 -*-
"""
目标系统导入文件生成：严格按模板列顺序 / 类型输出，异常日志随文件附带。

兼容性要点（对应验收项「格式兼容度」）：
- 列顺序、列名完全照抄目标模板，不多不少
- 无合并单元格、无公式、无多余索引列
- 数值写为数字、日期写为日期（或 yyyy-mm-dd 文本，可切换）
- 空值写空单元格，不写 NaN / None 字符串
"""
from __future__ import annotations

import math
from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .rules import RunResult, TargetColumn

HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")


def write_target_template(path: str, columns: list[TargetColumn], sheet_name: str = "导入模板"):
    """生成空白目标导入模板（仅表头）。"""
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        pd.DataFrame(columns=[c.label for c in columns]).to_excel(
            w, sheet_name=sheet_name, index=False)
        _style(w.sheets[sheet_name], columns)
    return path


def _style(ws, columns: list[TargetColumn] | None):
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    if columns:
        for i, c in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = c.width


def _coerce(value, dtype: str, date_mode: str):
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if dtype == "number":
        try:
            return float(value)
        except (TypeError, ValueError):
            return None
    if dtype == "date":
        if isinstance(value, datetime):
            d = value
        elif isinstance(value, pd.Timestamp):
            d = value.to_pydatetime()
        else:
            try:
                d = pd.to_datetime(value)
            except Exception:
                return None
        return d.strftime("%Y-%m-%d") if date_mode == "text" else d
    return str(value)


def export_result(result: RunResult, out_path: str,
                  target_columns: list[TargetColumn] | None = None,
                  date_mode: str = "date",
                  main_sheet: str = "导入数据") -> str:
    """把 RunResult 写成可直接导入目标系统的 xlsx。"""
    if target_columns:
        labels = [c.label for c in target_columns]
        dtypes = {c.label: c.dtype for c in target_columns}
        df = result.output.copy()
        # 输出列可能是标准字段 key，也可能是中文 label
        if all(c.key in df.columns for c in target_columns):
            df = df.rename(columns={c.key: c.label for c in target_columns})
        df = df.reindex(columns=labels)
    else:
        labels = list(result.output.columns)
        dtypes = {c: "text" for c in labels}
        df = result.output.copy()

    for label in labels:
        df[label] = df[label].map(lambda v: _coerce(v, dtypes.get(label, "text"), date_mode))
        if dtypes.get(label) == "date" and date_mode == "date":
            df[label] = pd.to_datetime(df[label], errors="coerce")

    with pd.ExcelWriter(out_path, engine="openpyxl") as w:
        df.to_excel(w, sheet_name=main_sheet, index=False)
        ws = w.sheets[main_sheet]
        _style(ws, target_columns)
        if date_mode == "date" and target_columns:
            for i, c in enumerate(target_columns, start=1):
                if c.dtype == "date":
                    for row in range(2, len(df) + 2):
                        ws.cell(row=row, column=i).number_format = "yyyy-mm-dd"

        res = result.anomaly_df()
        res.to_excel(w, sheet_name="异常日志", index=False)
        _style(w.sheets["异常日志"], None)
        w.sheets["异常日志"].column_dimensions["E"].width = 60

        if result.detail.get("settlement"):
            pd.DataFrame(result.detail["settlement"]).to_excel(
                w, sheet_name="计算明细", index=False)
            _style(w.sheets["计算明细"], None)

    return out_path
