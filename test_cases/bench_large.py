# -*- coding: utf-8 -*-
"""
超大规模基准测试（10 万行级）：生成数据 -> 引擎直测 -> 断言 -> 耗时报告。

场景：
  A. 10 万行请假单 + 2 万行余额表，假期余量结算（跨表逐单抵扣）
  B. 10 万行加班记录，部门 + 时长动态筛选

不依赖大模型（预置技能纯规则路径），直接调 core 引擎以隔离 Web 层噪声。
用法：python bench_large.py [行数]   （默认 100000）
"""
import random
import sys
import time
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from core.agent import TableBridgeAgent      # noqa: E402
from core.parser import load_table          # noqa: E402

TMP = ROOT / "data" / "bench"
TMP.mkdir(parents=True, exist_ok=True)
random.seed(2026)

DEPTS = ["生产部", "装配车间", "质检部", "仓储部", "销售部", "人事部", "财务部", "采购部"]


def gen_leave(n_rows: int) -> Path:
    """请假单：n_rows 行，含 0.3% 脏数据与少量空行。"""
    n_emp = max(1, n_rows // 5)
    p = TMP / "bench_请假单_10万.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["OA 系统请假明细导出"])
    ws.append([])
    ws.append(["员工编号", "姓名", "请假类型", "开始日期", "结束日期", "请假天数"])
    dirty = set(random.sample(range(n_rows), max(1, int(n_rows * 0.003))))
    blanks = {random.randrange(n_rows) for _ in range(5)}
    for i in range(n_rows):
        if i in blanks:
            ws.append([])     # 空行 -> 应进异常隔离
            continue
        emp = f"E{i % n_emp + 1:05d}"
        name = f"员工{i % n_emp + 1}"
        ltype = random.choice(["年假", "年假", "年假", "事假", "病假", "调休"])
        month = random.randint(1, 12)
        day = random.randint(1, 27)
        if i in dirty:
            ws.append([emp, name, ltype, f"2026/{month}/{day}", f"2026/{month}/{day + 2}", "???"])
            continue
        days = random.choice([1, 1, 2, 3, 5, 8, 12])
        ws.append([emp, name, ltype, f"2026-{month:02d}-{day:02d}",
                   f"2026-{month:02d}-{day:02d}", days])
    wb.save(p)
    return p


def gen_balance(n_emp: int) -> Path:
    p = TMP / "bench_年假余额_2万.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["工号", "姓名", "年假余额(天)"])
    for i in range(n_emp):
        ws.append([f"E{i + 1:05d}", f"员工{i + 1}", random.choice([0, 3, 5, 5, 8, 10, 15])])
    wb.save(p)
    return p


def gen_overtime(n_rows: int) -> Path:
    p = TMP / "bench_加班记录_10万.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ERP 加班记录导出"])
    ws.append([])
    ws.append(["工号", "姓名", "部门", "加班日期", "开始时间", "结束时间", "加班时长(小时)", "事由"])
    for i in range(n_rows):
        h = random.choice([1, 1.5, 2, 2, 2.5, 3, 4, 6])
        ws.append([f"E{i % 20000 + 1:05d}", f"员工{i % 20000 + 1}",
                   random.choice(DEPTS), f"2026-{random.randint(1, 12):02d}-{random.randint(1, 28):02d}",
                   "18:00", "21:00", h, "项目上线"])
    wb.save(p)
    return p


def report():
    n = int(sys.argv[1]) if len(sys.argv) > 1 else 100000
    n_emp = max(1, n // 5)
    print(f"== 超大规模基准：{n:,} 行 ==")

    # ---------- 场景 A：假期余量结算 ----------
    p_leave = gen_leave(n)
    p_bal = gen_balance(n_emp)
    t0 = time.perf_counter()
    leave_t = load_table(str(p_leave), "请假单_大表")
    t_parse_leave = time.perf_counter() - t0
    t0 = time.perf_counter()
    bal_t = load_table(str(p_bal), "余额表_大表")
    t_parse_bal = time.perf_counter() - t0
    print(f"[A] 解析请假单 {len(leave_t.df):,} 行：{t_parse_leave:.1f}s"
          f"（异常隔离 {len(leave_t.anomalies)} 条）；"
          f"解析余额表 {len(bal_t.df):,} 行：{t_parse_bal:.1f}s")

    t0 = time.perf_counter()
    out = TableBridgeAgent(llm=None).run("结算年假，先扣年假额度，超出转事假",
                                         [leave_t, bal_t], skill_key="leave_settlement")
    t_run = time.perf_counter() - t0
    res = out["result"]
    settlement = res.detail["settlement"]
    print(f"[A] 假期结算执行：{t_run:.1f}s，输出 {len(res.output):,} 行，"
          f"计算明细 {len(settlement):,} 条，异常 {len(res.anomalies):,} 条")

    # 不变量校验：抽样 500 名"无脏数据"员工，逐人复算 抵扣年假/转事假
    # （脏行/空行被隔离，直接按余额表全额可算的员工抽样验证）
    per_emp = {}
    for s in settlement:
        d = per_emp.setdefault(s["工号"], [0.0, 0.0])
        d[0] += s["抵扣年假"]
        d[1] += s["转事假"]
    bal_by_emp = {}
    for v in bal_t.df.itertuples(index=False):
        bal_by_emp[v.工号] = float(v._3 if hasattr(v, "_3") else v[2])
    # itertuples 列名带括号时用位置取值
    bal_by_emp = {}
    for row in bal_t.df.itertuples(index=False):
        bal_by_emp[row[0]] = float(row[2])
    # 只抽干净员工：请假单中无 "???" 行的工号集合
    # 注意：只有"年假"类型的单据才参与年假抵扣/转事假，事假/病假/调休直接透传
    dirty_emps = set()
    days_by_emp = {}
    for row in leave_t.df.itertuples(index=False):
        d = row[5]
        ok = d is not None and d == d
        if ok:
            if str(row[2]) == "年假":
                days_by_emp.setdefault(row[0], 0.0)
                days_by_emp[row[0]] += float(d)
        else:
            dirty_emps.add(row[0])
    sample = random.sample([e for e in days_by_emp if e not in dirty_emps
                            and e in bal_by_emp], min(500, len(days_by_emp)))
    bad = 0
    for e in sample:
        total_days = days_by_emp[e]
        bal = bal_by_emp[e]
        exp_deduct = min(total_days, bal)
        exp_personal = max(0.0, total_days - bal)
        got_deduct, got_personal = per_emp.get(e, [0.0, 0.0])
        if abs(got_deduct - exp_deduct) > 1e-6 or abs(got_personal - exp_personal) > 1e-6:
            bad += 1
            if bad <= 3:
                print(f"    MISMATCH {e}: exp=({exp_deduct},{exp_personal}) got=({got_deduct},{got_personal})")
    print(f"[A] 逐人复算抽样 {len(sample)} 人：{'全部一致 PASS' if bad == 0 else f'{bad} 人不一致 FAIL'}")

    # ---------- 场景 B：加班筛选 ----------
    p_ot = gen_overtime(n)
    t0 = time.perf_counter()
    ot_t = load_table(str(p_ot), "加班_大表")
    t_parse_ot = time.perf_counter() - t0
    print(f"[B] 解析加班记录 {len(ot_t.df):,} 行：{t_parse_ot:.1f}s")

    t0 = time.perf_counter()
    out_b = TableBridgeAgent(llm=None).run("筛选生产部加班时长大于等于2小时的记录",
                                           [ot_t], skill_key="overtime_import")
    t_run_b = time.perf_counter() - t0
    res_b = out_b["result"]
    print(f"[B] 筛选执行：{t_run_b:.1f}s，输出 {len(res_b.output):,} 行")

    # 期望行数直接从源表复算
    expect = 0
    for row in ot_t.df.itertuples(index=False):
        if row[2] == "生产部" and row[6] is not None and float(row[6]) >= 2:
            expect += 1
    ok_b = len(res_b.output) == expect
    print(f"[B] 期望 {expect:,} 行 / 实际 {len(res_b.output):,} 行：{'PASS' if ok_b else 'FAIL'}")

    total = t_parse_leave + t_parse_bal + t_run + t_parse_ot + t_run_b
    verdict = bad == 0 and ok_b
    print(f"== 总耗时 {total:.1f}s，基准{'通过' if verdict else '未通过'} ==")
    return 0 if verdict else 1


if __name__ == "__main__":
    sys.exit(report())
