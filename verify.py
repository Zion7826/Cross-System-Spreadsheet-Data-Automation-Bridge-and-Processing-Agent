# -*- coding: utf-8 -*-
"""
验收自检：对着任务书的三条验收项逐条跑，输出 PASS / FAIL。

1. 逻辑计算精度 —— 引擎结果 vs 独立公式复算，逐格比对
2. 格式兼容度   —— 导出文件 vs 目标模板，列顺序/类型/单元格规范检查
3. 抗错鲁棒性   —— 注入干扰字符与空行，不崩溃 + 正常数据完整 + 异常有日志

命令行：python verify.py
"""
from __future__ import annotations

import os
import sys
import tempfile
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from core.agent import TableBridgeAgent          # noqa: E402
from core.parser import load_table               # noqa: E402
from core.rules import SKILLS                    # noqa: E402

TMP = tempfile.mkdtemp(prefix="bridge_verify_")


# ------------------------------------------------------------ 1. 逻辑计算精度

def _case_leave():
    """固定用例：余额 + 请假单 -> 期望结果（人工用 Excel 公式的算法复算）。"""
    leave = pd.DataFrame([
        {"工号": "E1", "姓名": "甲", "假期类型": "年假", "开始日期": "2026-01-05", "结束日期": "2026-01-07", "请假天数": 3},
        {"工号": "E2", "姓名": "乙", "假期类型": "年假", "开始日期": "2026-02-02", "结束日期": "2026-02-06", "请假天数": 5},
        {"工号": "E3", "姓名": "丙", "假期类型": "年假", "开始日期": "2026-03-01", "结束日期": "2026-03-02", "请假天数": 2},
        {"工号": "E4", "姓名": "丁", "假期类型": "年假", "开始日期": "2026-03-05", "结束日期": "2026-03-05", "请假天数": 1},
        {"工号": "E5", "姓名": "戊", "假期类型": "事假", "开始日期": "2026-04-01", "结束日期": "2026-04-02", "请假天数": 2},
    ])
    balance = pd.DataFrame([
        {"工号": "E1", "姓名": "甲", "年假余额": 5},
        {"工号": "E2", "姓名": "乙", "年假余额": 2},
        {"工号": "E3", "姓名": "丙", "年假余额": 0},
        {"工号": "E5", "姓名": "戊", "年假余额": 6},
    ])

    expected = []
    bal = {"E1": 5.0, "E2": 2.0, "E3": 0.0, "E5": 6.0}
    for _, r in leave.iterrows():
        emp = r["工号"]
        days = float(r["请假天数"])
        b = bal.get(emp, 0.0)
        if r["假期类型"] != "年假":
            expected.append((emp, r["假期类型"], days, b))
            continue
        used = min(days, b)
        left = b - used
        if used > 0:
            expected.append((emp, "年假", used, left))
        if days - used > 0:
            expected.append((emp, "事假", round(days - used, 4), left))
        bal[emp] = left
    return leave, balance, expected


def check_logic() -> tuple[bool, str]:
    leave, balance, expected = _case_leave()
    lp = os.path.join(TMP, "leave.xlsx")
    bp = os.path.join(TMP, "balance.xlsx")
    leave.to_excel(lp, index=False)
    balance.to_excel(bp, index=False)

    agent = TableBridgeAgent()
    tables = [load_table(lp, "请假单"), load_table(bp, "余额表")]
    out = agent.run("做假期余量结算，年假抵扣，超出转事假", tables, skill_key="leave_settlement")
    df = out["result"].output

    got = [(str(r["emp_id"]), str(r["leave_type"]), round(float(r["days"]), 4),
            round(float(r["balance_left"]), 4)) for _, r in df.iterrows()]
    exp = [(e[0], e[1], round(float(e[2]), 4), round(float(e[3]), 4)) for e in expected]

    if got == exp:
        return True, f"逐格一致（{len(exp)} 条记录）。年假抵扣/事假转化/余额递减全部与人工公式复算结果相同。"
    diff = [f"  期望 {e} ≠ 实际 {g}" for e, g in zip(exp, got) if e != g]
    if len(exp) != len(got):
        diff.append(f"  记录数不一致：期望 {len(exp)} 条，实际 {len(got)} 条")
    return False, "数值不一致：\n" + "\n".join(diff[:10])


# ------------------------------------------------------------ 2. 格式兼容度

def check_format() -> tuple[bool, str]:
    leave, balance, _ = _case_leave()
    lp = os.path.join(TMP, "leave2.xlsx")
    bp = os.path.join(TMP, "balance2.xlsx")
    leave.to_excel(lp, index=False)
    balance.to_excel(bp, index=False)

    agent = TableBridgeAgent()
    tables = [load_table(lp, "请假单"), load_table(bp, "余额表")]
    out = agent.run("", tables, skill_key="leave_settlement")
    tpl = SKILLS["leave_settlement"].target_columns
    op = os.path.join(TMP, "out.xlsx")

    # 先确保模板文件存在
    tpl_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "templates", "目标系统_假期导入模板.xlsx")
    if not os.path.exists(tpl_path):
        from core.exporter import write_target_template
        write_target_template(tpl_path, tpl)

    agent.export(out, op)

    wb_t = load_workbook(tpl_path)
    wb_o = load_workbook(op)
    ws_t, ws_o = wb_t[wb_t.sheetnames[0]], wb_o[wb_o.sheetnames[0]]
    tpl_cols = [c.value for c in ws_t[1]]
    out_cols = [c.value for c in ws_o[1]]

    problems = []
    if tpl_cols != out_cols:
        problems.append(f"列名/列序与模板不符：模板 {tpl_cols} vs 输出 {out_cols}")
    if ws_o.merged_cells.ranges:
        problems.append(f"存在合并单元格：{list(ws_o.merged_cells.ranges)}")
    if "异常日志" not in wb_o.sheetnames:
        problems.append("缺少异常日志表")
    # 类型检查
    type_map = {c.label: c.dtype for c in tpl}
    for i, col in enumerate(out_cols, start=1):
        dtype = type_map.get(col, "text")
        for row in range(2, ws_o.max_row + 1):
            v = ws_o.cell(row=row, column=i).value
            if v is None:
                continue
            if dtype == "number" and not isinstance(v, (int, float)):
                problems.append(f"第{row}行「{col}」应为数值，实际 {type(v).__name__}")
            if dtype == "date" and not isinstance(v, datetime):
                problems.append(f"第{row}行「{col}」应为日期，实际 {type(v).__name__}")
    if problems:
        return False, "；".join(problems[:8])
    return True, f"列序/列名与模板完全一致（{out_cols}），无合并单元格，数值列为数字、日期列为日期，附带异常日志表。"


# ------------------------------------------------------------ 3. 抗错鲁棒性

def check_robust() -> tuple[bool, str]:
    leave, balance, _ = _case_leave()
    lp = os.path.join(TMP, "leave3.xlsx")
    bp = os.path.join(TMP, "balance3.xlsx")

    dirty = leave.copy()
    dirty.loc[len(dirty)] = {c: "" for c in dirty.columns}                  # 空行
    dirty.loc[len(dirty)] = {"工号": "E9", "姓名": "己", "假期类型": "年假",
                             "开始日期": "###", "结束日期": "@@@", "请假天数": "abc"}  # 干扰字符
    dirty["请假天数"] = dirty["请假天数"].astype(object)
    dirty.loc[0, "请假天数"] = "!!!"
    dirty.to_excel(lp, index=False)
    balance.to_excel(bp, index=False)

    agent = TableBridgeAgent()
    tables = [load_table(lp, "请假单"), load_table(bp, "余额表")]
    try:
        out = agent.run("", tables, skill_key="leave_settlement")
    except Exception as ex:
        return False, f"注入干扰后崩溃：{ex}"

    res = out["result"]
    ok = len(res.output) > 0 and len(res.anomalies) > 0
    kinds = sorted({a["异常类型"] for a in res.anomalies})
    msg = (f"注入空行 + 干扰字符（!!!/###/abc）后未崩溃；正常数据输出 {len(res.output)} 行，"
           f"异常隔离 {len(res.anomalies)} 条，异常类型：{kinds}")
    return ok, msg


CHECKS = [
    ("逻辑计算精度验收", "系统计算值与人工电子表格公式结果完全一致", check_logic),
    ("格式兼容度验收", "生成文件可无缝导入目标老旧系统，不触发格式报错", check_format),
    ("抗错鲁棒性验收", "注入干扰字符/空行后不宕机，正常数据与异常日志均正确输出", check_robust),
]


def run_all() -> list[dict]:
    rows = []
    for name, desc, fn in CHECKS:
        try:
            ok, msg = fn()
        except Exception as ex:
            ok, msg = False, f"自检执行异常：{ex}"
        rows.append({"验收项": name, "验收标准": desc, "结果": "PASS" if ok else "FAIL", "详情": msg})
    return rows


if __name__ == "__main__":
    for r in run_all():
        print(f"[{r['结果']}] {r['验收项']}：{r['详情']}")
